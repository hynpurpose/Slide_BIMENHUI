import json
from pathlib import Path

import openpyxl

wb = openpyxl.load_workbook('创维词条分类.xlsx', read_only=True)


def sheet_to_rows(ws):
    headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not any(c is not None for c in row):
            continue
        d = {}
        for h, v in zip(headers, row):
            d[h] = v if v is not None else ''
        rows.append(d)
    return rows


data = {
    'meta': {
        'sourceFile': '创维词条分类.xlsx',
    },
    'optimization': sheet_to_rows(wb['优化词']),
    'monitor': sheet_to_rows(wb['监测词']),
}
data['meta']['optimizationCount'] = len(data['optimization'])
data['meta']['monitorCount'] = len(data['monitor'])

out = Path('src/data/skyworthKeywords.json')
out.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'Wrote {out}')
print(f'optimization: {data["meta"]["optimizationCount"]}, monitor: {data["meta"]["monitorCount"]}')
